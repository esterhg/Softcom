# -*- coding: utf-8 -*-
"""
Carga (idempotente) las rutinas de mantenimiento de ascensores a partir del
protocolo anual en Excel.

Uso:
    python manage.py seed_ascensores
    python manage.py seed_ascensores --file "PROTOCOLO MENSUAL ASCENSORES ....xlsx"
    python manage.py seed_ascensores --dry-run

El protocolo es un PLAN ANUAL ROTATIVO: cada mes toca un protocolo distinto.
Cada hoja del Excel = un protocolo = una Rutina. Dentro de cada hoja hay 3 niveles:
    ZONA        -> paso HEADER (encabezado de sección, ej. FOSO, CABINA SUPERIOR)
    COMPONENTE  -> paso HEADER secundario (ej. Panel luminoso)
    ACTIVIDAD   -> paso CHECK (o NUMERICO si pide medir/anotar un valor)
"""
import os
from datetime import timedelta

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction


# ──────────────────────────────────────────────────────────────────────────
# Definición de los protocolos (hoja de Excel -> metadatos de la Rutina)
# meses: en qué meses del año se ejecuta este protocolo (plan rotativo)
# ──────────────────────────────────────────────────────────────────────────
PROTOCOLOS = [
    {"hoja": "LIMPIEZA (ENE,MAY,SEP)",  "codigo": "ASC-LIMPIEZA",   "nombre": "Protocolo LIMPIEZA - Ascensores",      "meses": [1, 5, 9],   "horas": 2},
    {"hoja": "SEGURIDAD I (FEB)",       "codigo": "ASC-SEGURIDAD-I", "nombre": "Protocolo SEGURIDAD I - Ascensores",   "meses": [2],         "horas": 3},
    {"hoja": "PUERTA (MAR, JUL, NOV)",  "codigo": "ASC-PUERTA",     "nombre": "Protocolo PUERTA - Ascensores",        "meses": [3, 7, 11],  "horas": 2},
    {"hoja": "COMANDO (ABR,OCT)",       "codigo": "ASC-COMANDO",    "nombre": "Protocolo COMANDO - Ascensores",       "meses": [4, 10],     "horas": 3},
    {"hoja": "MAQUINARIA (JUN,DEC)",    "codigo": "ASC-MAQUINARIA", "nombre": "Protocolo MAQUINARIA - Ascensores",    "meses": [6, 12],     "horas": 3},
    {"hoja": "SEGURIDAD II (AGO)",      "codigo": "ASC-SEGURIDAD-II", "nombre": "Protocolo SEGURIDAD II - Ascensores", "meses": [8],        "horas": 3},
]

DEFAULT_FILE = "PROTOCOLO MENSUAL ASCENSORES 01.01.2025 1 1 3 1 .xlsx"

# Zonas físicas conocidas del ascensor (nivel 1 = HEADER principal)
ZONAS = {
    "cabina interna", "subtecho y ventilador", "piso", "maquinario y comando",
    "cabina superior", "caja de carrera", "foso", "estructura de la cabina",
}

# Verbos / prefijos que identifican una ACTIVIDAD ejecutable por el técnico
VERBOS = (
    "compruebe", "comprueba", "comprobar", "verifique", "verificar", "verifque",
    "inspeccione", "inspeccionar", "inspecione", "revise", "revisar", "revise y",
    "mida", "medir", "prueba", "pruebe", "preuba", "prube", "testear", "teste",
    "limpieza", "ajuste", "active", "existencia", "funcionamiento",
    "integridad del", "integridad de los", "integridad de la", "integridad de las",
    "para los", "cuando el", "en los modelos", "nivelación de", "comprobación de",
    "pruebe", "preuba el", "integridad,",
)

# Pistas de que la actividad captura un VALOR numérico (medición a anotar)
MED_HINTS = (
    "ha ", "hb ", "hc ", "lcab", "lcp", "750 mm", "150mm", "diametro",
    "diámetro", "amperaje", "abertura", "25mm", "altura mínima", "altura minima",
)


def _is_actividad(txt: str) -> bool:
    low = txt.lower()
    if any(low.startswith(v) for v in VERBOS):
        return True
    if low.startswith(("lcab", "lcp", "mida")):
        return True
    return False


def _is_numerico(txt: str) -> bool:
    low = txt.lower()
    return any(h in low for h in MED_HINTS)


def _parse_sheet(ws):
    """Devuelve lista de (nivel, texto) donde nivel ∈ {ZONA, COMPONENTE, ACTIVIDAD, MEDICION}."""
    out = []
    for r in range(2, ws.max_row + 1):  # fila 1 = título "PROTOCOLO X"
        raw = ws.cell(row=r, column=1).value
        if raw is None or str(raw).strip() == "":
            continue
        # Una celda puede traer varias líneas (ej. HA/HB/HC)
        for part in str(raw).split("\n"):
            txt = part.strip()
            if not txt:
                continue
            low = txt.lower()
            if low in ZONAS:
                out.append(("ZONA", txt))
            elif _is_actividad(txt):
                out.append(("MEDICION" if _is_numerico(txt) else "ACTIVIDAD", txt))
            else:
                out.append(("COMPONENTE", txt))
    return out


class Command(BaseCommand):
    help = "Carga las rutinas de mantenimiento de ascensores desde el protocolo Excel (idempotente)."

    def add_arguments(self, parser):
        parser.add_argument("--file", type=str, default=None, help="Ruta al Excel del protocolo.")
        parser.add_argument("--dry-run", action="store_true", help="No escribe en la BD; solo muestra el resumen.")

    def handle(self, *args, **opts):
        try:
            import openpyxl  # noqa
        except ImportError:
            raise CommandError("openpyxl no está instalado. Ejecuta: pip install openpyxl")

        file_arg = opts["file"] or DEFAULT_FILE
        path = file_arg if os.path.isabs(file_arg) else os.path.join(settings.BASE_DIR, file_arg)
        if not os.path.exists(path):
            raise CommandError(f"No se encontró el archivo: {path}")

        dry = opts["dry_run"]

        from activos.models import Categoria
        from mantenimiento.models import Rutina, PasoRutina, Frecuencia, Tipo

        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)

        self.stdout.write(self.style.MIGRATE_HEADING(f"Protocolo: {path}"))

        with transaction.atomic():
            # 1) Categoría de activo "Ascensores" (reutiliza una existente, sin duplicar)
            categoria_activo = Categoria.objects.filter(nombre__iexact="Ascensores").first()
            if not categoria_activo:
                categoria_activo = Categoria.objects.create(nombre="Ascensores")
                self.stdout.write(self.style.SUCCESS("  ✓ Categoría de activo 'Ascensores' creada."))
            else:
                self.stdout.write(f"  • Categoría de activo existente reutilizada: '{categoria_activo.nombre}' (id={categoria_activo.id}).")

            # 2) Tipo de mantenimiento raíz "Ascensores"
            tipo_root, _ = Tipo.objects.get_or_create(
                codigo="ASC",
                defaults={"nombre": "Ascensores", "categoria_activo": categoria_activo},
            )
            if tipo_root.categoria_activo_id != categoria_activo.id:
                tipo_root.categoria_activo = categoria_activo
                tipo_root.save(update_fields=["categoria_activo"])

            # 3) Frecuencia "Mensual" (el plan corre todos los meses, protocolo rotativo)
            frec_mensual, _ = Frecuencia.objects.get_or_create(
                nombre="Mensual", defaults={"dias": 30}
            )

            total_rutinas = 0
            total_pasos = 0

            for proto in PROTOCOLOS:
                hoja = proto["hoja"]
                if hoja not in wb.sheetnames:
                    self.stdout.write(self.style.WARNING(f"  ⚠ Hoja no encontrada: {hoja} (se omite)"))
                    continue

                filas = _parse_sheet(wb[hoja])
                nombre_corto = proto["nombre"].split(" - ")[0].replace("Protocolo ", "")

                # Tipo hijo por protocolo
                tipo_proto, _ = Tipo.objects.get_or_create(
                    codigo=proto["codigo"] + "-TIPO",
                    defaults={"nombre": nombre_corto, "padre": tipo_root,
                              "categoria_activo": categoria_activo},
                )

                meses_txt = ", ".join(
                    ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"][m - 1]
                    for m in proto["meses"]
                )
                descripcion = f"Protocolo de ascensores ejecutado en: {meses_txt}."

                if dry:
                    n_act = sum(1 for n, _ in filas if n in ("ACTIVIDAD", "MEDICION"))
                    n_hdr = sum(1 for n, _ in filas if n in ("ZONA", "COMPONENTE"))
                    self.stdout.write(f"  {proto['codigo']:18} {nombre_corto:14} secciones={n_hdr:3} actividades={n_act:3}  ({meses_txt})")
                    continue

                # Rutina (idempotente por codigo_rutina)
                rutina, created = Rutina.objects.update_or_create(
                    codigo_rutina=proto["codigo"],
                    defaults={
                        "nombre": proto["nombre"],
                        "categoria_activo": categoria_activo,
                        "tipo": tipo_proto,
                        "frecuencia": frec_mensual,
                        "descripcion": descripcion,
                        "tiempo_estimado": timedelta(hours=proto["horas"]),
                        "cantidad_tecnicos": 1,
                    },
                )

                # Reconstruir los pasos desde cero (para que re-ejecutar refleje el Excel)
                rutina.pasos.all().delete()

                orden = 0
                pasos_bulk = []
                for nivel, texto in filas:
                    if nivel == "ZONA":
                        tipo_resp = "HEADER"
                    elif nivel == "COMPONENTE":
                        tipo_resp = "HEADER"
                    elif nivel == "MEDICION":
                        tipo_resp = "NUMERICO"
                    else:  # ACTIVIDAD
                        tipo_resp = "CHECK"

                    pasos_bulk.append(PasoRutina(
                        rutina=rutina,
                        orden=orden,
                        descripcion=texto,
                        tipo_respuesta=tipo_resp,
                    ))
                    orden += 1

                PasoRutina.objects.bulk_create(pasos_bulk)
                total_rutinas += 1
                total_pasos += len(pasos_bulk)

                estado = "creada" if created else "actualizada"
                self.stdout.write(self.style.SUCCESS(
                    f"  ✓ {proto['codigo']:18} {estado:12} ({len(pasos_bulk)} pasos)  meses: {meses_txt}"
                ))

            if dry:
                self.stdout.write(self.style.WARNING("DRY-RUN: no se escribió nada en la base de datos."))
                transaction.set_rollback(True)
                return

        self.stdout.write(self.style.MIGRATE_HEADING(
            f"\nListo. {total_rutinas} rutinas, {total_pasos} pasos. Categoría de activo: 'Ascensores'."
        ))
